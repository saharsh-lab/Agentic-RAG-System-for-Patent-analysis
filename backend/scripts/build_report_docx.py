"""Assemble docs/report/*.md into one Word document (docs/report/Project_Report.docx).

    cd backend && .venv/bin/python scripts/build_report_docx.py

Re-run after any change to the chapters. The output is a clean draft to paste into the
college template: headings, paragraphs, lists, tables, figures and references. It handles
the Markdown used in these chapters, not Markdown in general.
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / "docs" / "report"
DIAGRAMS = ROOT / "docs" / "diagrams"
CHAPTERS = [
    "01_abstract.md",
    "02_introduction.md",
    "03_literature_review.md",
    "04_system_design.md",
    "05_implementation.md",
    "06_evaluation_methodology.md",
    "07_results.md",
    "08_discussion.md",
    "09_conclusion.md",
    "references.md",
]
TITLE = (
    "Agentic RAG-Based Patent Intelligence System with Live Data Retrieval "
    "and Hallucination Detection"
)
# Notes to the authors, not part of the report
AUTHOR_NOTE = re.compile(
    r"^> (All references were checked|Checked on |\*\*Template|Draft\.|Revise )"
)
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
FIGURE = re.compile(r"docs/diagrams/(\w+)")


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\((?:https?://|\.\.?/)[^)]+\)", r"\1", text)  # links → text
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, row[j] if j < len(row) else "")
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_figure(doc, caption_line: str) -> None:
    for name in FIGURE.findall(caption_line):
        image = DIAGRAMS / f"{name}.png"
        if image.exists():
            doc.add_picture(str(image), width=Inches(6.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_chapter(doc, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if AUTHOR_NOTE.match(line):
            while i < len(lines) and lines[i].startswith(">"):
                i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith(">"):
            text = []
            while i < len(lines) and lines[i].startswith(">"):
                text.append(lines[i].lstrip("> ").strip())
                i += 1
            add_inline(doc.add_paragraph(style="Quote"), " ".join(text))
            continue
        if m := re.match(r"^(#{1,3}) (.*)", line):
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
        elif m := re.match(r"^\s*[-*] (.*)", line):
            text = [m.group(1)]
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]):
                i += 1
                text.append(lines[i].strip())
            add_inline(doc.add_paragraph(style="List Bullet"), " ".join(text))
        elif m := re.match(r"^\d+\. (.*)", line):
            text = [m.group(1)]
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]):
                i += 1
                text.append(lines[i].strip())
            add_inline(doc.add_paragraph(style="List Number"), " ".join(text))
        elif line.strip():
            text = [line.strip()]
            while (
                i + 1 < len(lines)
                and lines[i + 1].strip()
                and not re.match(r"^(#|\||>|\s*[-*] |\d+\. )", lines[i + 1])
            ):
                i += 1
                text.append(lines[i].strip())
            paragraph_text = " ".join(text)
            if FIGURE.search(paragraph_text) and paragraph_text.startswith("*Figure"):
                add_figure(doc, paragraph_text)
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(caption, paragraph_text)
            else:
                add_inline(doc.add_paragraph(), paragraph_text)
        i += 1


def main() -> int:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    for text, size, bold in (
        (TITLE, 20, True),
        ("B.Tech Major Project Report, Computer Science and Engineering", 13, False),
        ("Guide: Dr. T. Swathi", 13, False),
        ("[Team members and roll numbers]", 12, False),
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
    doc.add_page_break()

    for name in CHAPTERS:
        add_chapter(doc, REPORT / name)
        if name != CHAPTERS[-1]:
            doc.add_page_break()

    out = REPORT / "Project_Report.docx"
    doc.save(out)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
