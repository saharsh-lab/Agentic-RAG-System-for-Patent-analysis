"""Builders for test documents, so no binary files need to be committed."""

import io
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"


def fixture_bytes(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def make_pdf(pages: list[list[str]], header: str | None = None) -> bytes:
    """A PDF with one text box per block, an optional running header, and page numbers."""
    import pymupdf

    doc = pymupdf.open()
    for number, blocks in enumerate(pages, start=1):
        page = doc.new_page()
        if header:
            page.insert_text((72, 40), header.format(page=number), fontsize=9)
        y = 80
        for block in blocks:
            height = 30 + 14 * (len(block) // 70 + block.count("\n"))
            page.insert_textbox(pymupdf.Rect(72, y, 520, y + height), block, fontsize=10)
            y += height + 10
        page.insert_text((300, 810), str(number), fontsize=9)
    return doc.tobytes()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
