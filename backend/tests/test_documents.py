import io
import fitz
import docx
import pytest
from httpx import AsyncClient, ASGITransport
from app.main import app as fastapi_app
from app.documents.parser import DocumentParser
from app.documents.chunker import PatentDocumentChunker


def create_sample_pdf_bytes() -> bytes:
    doc = fitz.open()
    
    # Page 1: Abstract & Background
    page1 = doc.new_page()
    page1.insert_text((50, 50), "ABSTRACT\nThis invention relates to a solid-state electrolyte composition for lithium batteries.")
    page1.insert_text((50, 150), "BACKGROUND OF THE INVENTION\nLithium batteries often suffer from thermal runaway. Prior art liquid electrolytes are flammable.")

    # Page 2: Summary & Detailed Description
    page2 = doc.new_page()
    page2.insert_text((50, 50), "SUMMARY OF THE INVENTION\nThe present invention provides a sulfide-based electrolyte that exhibits high ionic conductivity.")
    page2.insert_text((50, 150), "DETAILED DESCRIPTION\nThe sulfide electrolyte comprises Li2S and P2S5 in a ratio ranging from 70:30 to 80:20.")

    # Page 3: Claims
    page3 = doc.new_page()
    page3.insert_text((50, 50), "CLAIMS\n1. A solid-state lithium battery comprising a sulfide electrolyte.\n2. The battery of claim 1 wherein the ratio is 75:25.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("ABSTRACT", level=1)
    doc.add_paragraph("A high-energy density solar panel with multi-junction photovoltaic cells.")
    doc.add_heading("CLAIMS", level=1)
    doc.add_paragraph("1. A solar panel comprising multi-junction cells.")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


@pytest.mark.asyncio
async def test_pdf_parsing():
    pdf_bytes = create_sample_pdf_bytes()
    parsed = DocumentParser.parse(pdf_bytes, "sample_patent.pdf")
    
    assert parsed.filename == "sample_patent.pdf"
    assert parsed.file_type == "pdf"
    assert len(parsed.pages) == 3
    
    # Page 1 checks
    assert parsed.pages[0].page_number == 1
    assert "ABSTRACT" in parsed.pages[0].text
    assert parsed.pages[0].section in ["Abstract", "Background"]
    
    # Page 3 claims checks
    assert parsed.pages[2].page_number == 3
    assert "CLAIMS" in parsed.pages[2].text
    assert parsed.pages[2].section == "Claims"


@pytest.mark.asyncio
async def test_chunker_page_and_section_preservation():
    pdf_bytes = create_sample_pdf_bytes()
    parsed = DocumentParser.parse(pdf_bytes, "sample_patent.pdf")
    chunker = PatentDocumentChunker(chunk_size=300, chunk_overlap=50)
    chunks = chunker.chunk_document(parsed)
    
    assert len(chunks) > 0
    # Every chunk must preserve page_number >= 1
    for chunk in chunks:
        assert chunk.page_number in [1, 2, 3]
        assert chunk.section is not None
        assert len(chunk.content) > 0


@pytest.mark.asyncio
async def test_document_upload_and_management_api():
    pdf_bytes = create_sample_pdf_bytes()
    transport = ASGITransport(app=fastapi_app)
    
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Upload PDF document
        files = {"file": ("test_patent.pdf", pdf_bytes, "application/pdf")}
        response = await client.post("/api/v1/documents/upload", files=files)
        assert response.status_code == 201
        data = response.json()
        doc_id = data["id"]
        assert data["filename"] == "test_patent.pdf"
        assert data["file_type"] == "pdf"
        assert data["chunk_count"] > 0
        
        # 2. List documents
        list_resp = await client.get("/api/v1/documents/")
        assert list_resp.status_code == 200
        docs = list_resp.json()
        assert any(d["id"] == doc_id for d in docs)
        
        # 3. Get document details with chunks
        detail_resp = await client.get(f"/api/v1/documents/{doc_id}")
        assert detail_resp.status_code == 200
        detail = detail_resp.json()
        assert detail["id"] == doc_id
        assert len(detail["chunks"]) > 0
        # Check chunk page numbers are preserved
        assert detail["chunks"][0]["page_number"] is not None
        
        # 4. Delete document
        del_resp = await client.delete(f"/api/v1/documents/{doc_id}")
        assert del_resp.status_code == 204
        
        # 5. Verify deletion
        get_again = await client.get(f"/api/v1/documents/{doc_id}")
        assert get_again.status_code == 404
