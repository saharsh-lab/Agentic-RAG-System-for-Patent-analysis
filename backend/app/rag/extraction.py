"""Step 1 of ingestion: check what a file really is, then pull text out of it.

We never trust the filename extension alone. Every format has a recognisable
"magic number" at the start of the file (PDF files start with `%PDF-`, DOCX
files are ZIP archives starting with `PK`). A file called `report.pdf` that is
really an executable is rejected here.

Output is a list of pages, each a list of text *blocks* (roughly paragraphs).
Keeping page numbers is what later lets an answer cite "Patent A, page 4".
"""

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Literal

from app.core.errors import (
    EmptyDocumentError,
    InvalidDocumentError,
    UnsupportedFileError,
)

FileKind = Literal["pdf", "docx", "txt"]

MIME_TYPES: dict[FileKind, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}

MIN_TEXT_CHARS = 20


@dataclass
class Page:
    number: int | None  # 1-based PDF page; None for formats without pages (DOCX, TXT)
    blocks: list[str] = field(default_factory=list)


def detect_file_kind(filename: str, data: bytes) -> FileKind:
    """Decide the file type from its extension AND its content; reject mismatches."""
    extension = PurePath(filename).suffix.lower()
    if extension == ".pdf":
        if not data.lstrip()[:5] == b"%PDF-":
            raise InvalidDocumentError("This file has a .pdf name but is not a valid PDF.")
        return "pdf"
    if extension == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise InvalidDocumentError("This file has a .docx name but is not a valid DOCX.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise InvalidDocumentError("This archive is not a Word document.")
                _check_zip_bomb(archive)
        except zipfile.BadZipFile as exc:
            raise InvalidDocumentError("This DOCX file is corrupted.") from exc
        return "docx"
    if extension == ".txt":
        if b"\x00" in data[:8192]:
            raise InvalidDocumentError("This .txt file contains binary data, not text.")
        return "txt"
    raise UnsupportedFileError(
        f"Unsupported file type '{extension or 'none'}'. Upload a PDF, DOCX or TXT file."
    )


# A DOCX is a ZIP archive. A "zip bomb" is a small file that expands to gigabytes when
# unpacked; reject archives whose contents are implausibly large for a document.
MAX_DOCX_UNCOMPRESSED = 200 * 1024 * 1024
MAX_DOCX_ENTRIES = 5000
MAX_COMPRESSION_RATIO = 200


def _check_zip_bomb(archive: zipfile.ZipFile) -> None:
    entries = archive.infolist()
    total = sum(e.file_size for e in entries)
    compressed = sum(e.compress_size for e in entries) or 1
    if (
        len(entries) > MAX_DOCX_ENTRIES
        or total > MAX_DOCX_UNCOMPRESSED
        or total / compressed > MAX_COMPRESSION_RATIO
    ):
        raise InvalidDocumentError(
            "This DOCX file expands to an implausible size and was rejected."
        )


def extract_pages(kind: FileKind, data: bytes, *, max_pages: int = 500) -> list[Page]:
    if kind == "pdf":
        pages = _extract_pdf(data, max_pages=max_pages)
    elif kind == "docx":
        pages = _extract_docx(data)
    else:
        pages = _extract_txt(data)

    total_chars = sum(len(block) for page in pages for block in page.blocks)
    if total_chars < MIN_TEXT_CHARS:
        hint = " It may be a scanned image; OCR is not supported yet." if kind == "pdf" else ""
        raise EmptyDocumentError(f"No readable text was found in this document.{hint}")
    return pages


def _extract_pdf(data: bytes, *, max_pages: int) -> list[Page]:
    import pymupdf  # imported lazily: only needed for PDFs

    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - PyMuPDF raises several error types
        raise InvalidDocumentError("This PDF is damaged and could not be opened.") from exc

    with document:
        if document.needs_pass:
            raise InvalidDocumentError("This PDF is password-protected.")
        if document.page_count > max_pages:
            raise InvalidDocumentError(
                f"This PDF has {document.page_count} pages; the limit is {max_pages}."
            )
        pages = []
        for index, pdf_page in enumerate(document):
            # "blocks" groups lines into paragraphs; sort=True gives natural reading order.
            # Each item is (x0, y0, x1, y1, text, block_no, block_type); type 0 = text.
            blocks = [
                item[4].strip()
                for item in pdf_page.get_text("blocks", sort=True)
                if item[6] == 0 and item[4].strip()
            ]
            pages.append(Page(number=index + 1, blocks=blocks))
    return pages


def _extract_docx(data: bytes) -> list[Page]:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise InvalidDocumentError("This Word document could not be read.") from exc

    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # Tables are common in patents (e.g. experimental results); keep them as text rows
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))  # merged cells repeat
    return [Page(number=None, blocks=blocks)]


def _extract_txt(data: bytes) -> list[Page]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise InvalidDocumentError("This text file's encoding is not supported (use UTF-8).")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
    return [Page(number=None, blocks=blocks)]
